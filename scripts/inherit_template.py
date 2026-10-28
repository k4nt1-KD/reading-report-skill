# -*- coding: utf-8 -*-
"""Reuse an approved deck as the template base.

Why: after a few review rounds the original .potx/.pptx template is often gone,
but the last approved deck still carries the master, layouts, header lockup and
footer branding. Dropping its slides keeps all of that and gives a clean canvas.

    python inherit_template.py --base old.pptx --out build/base.pptx --list

Then build with ``deck_kit.Deck("build/base.pptx")``.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

IN = 914400

if hasattr(sys.stdout, "reconfigure"):  # Windows consoles default to GBK
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def describe(prs: Presentation) -> None:
    print(f"slides={len(prs.slides)}  size={prs.slide_width / IN:.3f} x {prs.slide_height / IN:.3f} in")
    for mi, master in enumerate(prs.slide_masters):
        print(f"master {mi}: {len(master.shapes)} shapes, theme={master.theme.name if hasattr(master, 'theme') else '?'}")
    for li, layout in enumerate(prs.slide_layouts):
        print(f"layout {li}: {layout.name!r}")
        for sh in layout.shapes:
            txt = sh.text_frame.text.strip()[:24] if sh.has_text_frame else ""
            print(f"    [{sh.shape_type}] {sh.name!r} "
                  f"({sh.left / IN:.2f},{sh.top / IN:.2f},{sh.width / IN:.2f},{sh.height / IN:.2f}) {txt!r}")
        for ph in layout.placeholders:
            print(f"    ph idx={ph.placeholder_format.idx} name={ph.name!r} "
                  f"({ph.left / IN:.2f},{ph.top / IN:.2f})")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True, help="deck whose branding should be inherited")
    ap.add_argument("--out", help="where to write the stripped copy")
    ap.add_argument("--list", action="store_true", help="print layouts / placeholders / master shapes")
    args = ap.parse_args()

    prs = Presentation(str(Path(args.base).resolve()))
    if args.list:
        describe(prs)

    if not args.out:
        return

    before = len(prs.slides)
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)

    out = Path(args.out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"[inherit] dropped {before} slides, kept {len(prs.slide_layouts)} layouts "
          f"and {len(prs.slide_masters)} master(s) -> {out}")


if __name__ == "__main__":
    main()

